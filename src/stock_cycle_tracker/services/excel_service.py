"""Excel import, parsing, validation, and analytical export service."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
import io
from pathlib import Path
import re
from typing import Any, List, Optional, Sequence, Tuple, Union
import urllib.request

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from stock_cycle_tracker.domain.models import CycleAnalysis
from stock_cycle_tracker.services.cycle_service import CycleService


@dataclass
class ExcelRowData:
    row_index: int
    stock_name: str
    reference_date: date
    raw_date_str: str


@dataclass
class ExcelRowError:
    row_index: int
    raw_stock: str
    raw_date: str
    error_message: str


@dataclass
class ExcelImportResult:
    total_rows: int = 0
    valid_rows: List[ExcelRowData] = field(default_factory=list)
    invalid_rows: List[ExcelRowError] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


def convert_google_sheets_url(url: str) -> str:
    """Converts a public Google Sheets / Google Docs share URL into a direct XLSX download URL."""
    url = url.strip()
    match = re.search(r'/spreadsheets/d/([a-zA-Z0-9-_]+)', url)
    if match:
        doc_id = match.group(1)
        gid_match = re.search(r'[#&?]gid=([0-9]+)', url)
        export_url = f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=xlsx"
        if gid_match:
            export_url += f"&gid={gid_match.group(1)}"
        return export_url
    return url


class ExcelService:
    """Handles Excel file ingestion, row validation, URL downloads, and 18-column analytical exports."""

    STOCK_COL_VARIANTS = ["stock name", "stock", "stock symbol", "symbol", "instrument"]
    DATE_COL_VARIANTS = ["reference date", "date", "ld", "last date", "ref date", "research date"]

    def _parse_date(self, val: Any) -> Optional[date]:
        if val is None or pd.isna(val):
            return None
        if isinstance(val, (datetime, pd.Timestamp)):
            return val.date()
        if isinstance(val, date):
            return val

        val_str = str(val).strip()
        if not val_str:
            return None

        # Common format attempts
        formats = [
            "%d-%b-%Y", "%d-%b-%y", "%d-%B-%Y", "%d-%m-%Y", "%Y-%m-%d",
            "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d"
        ]
        for fmt in formats:
            try:
                return datetime.strptime(val_str, fmt).date()
            except ValueError:
                pass

        try:
            parsed = pd.to_datetime(val_str, errors="coerce")
            if pd.notna(parsed):
                return parsed.date()
        except Exception:
            pass

        return None

    def validate_and_parse_url(self, raw_url: str) -> ExcelImportResult:
        """Fetches an Excel file from a Google Sheets public link or HTTP URL and parses it."""
        result = ExcelImportResult()
        clean_url = convert_google_sheets_url(raw_url.strip())

        try:
            req = urllib.request.Request(
                clean_url,
                headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"},
            )
            with urllib.request.urlopen(req, timeout=15) as resp:
                content_bytes = resp.read()
                return self.validate_and_parse_upload(content_bytes)
        except Exception as e:
            result.warnings.append(f"Failed to download Excel file from URL ({e}). Ensure the link has public view access.")
            return result

    def validate_and_parse_upload(
        self,
        file_input: Union[str, Path, bytes, io.BytesIO],
    ) -> ExcelImportResult:
        result = ExcelImportResult()

        try:
            if isinstance(file_input, (str, Path)):
                # If it's a URL string
                if str(file_input).startswith("http://") or str(file_input).startswith("https://"):
                    return self.validate_and_parse_url(str(file_input))
                df = pd.read_excel(file_input)
            elif isinstance(file_input, bytes):
                df = pd.read_excel(io.BytesIO(file_input))
            else:
                df = pd.read_excel(file_input)
        except Exception as e:
            result.warnings.append(f"Failed to read Excel file: {e}")
            return result

        if df.empty:
            result.warnings.append("The uploaded Excel sheet is empty.")
            return result

        # Identify required columns
        stock_col = None
        date_col = None

        col_map = {str(c).strip().lower(): c for c in df.columns}
        for variant in self.STOCK_COL_VARIANTS:
            if variant in col_map:
                stock_col = col_map[variant]
                break

        for variant in self.DATE_COL_VARIANTS:
            if variant in col_map:
                date_col = col_map[variant]
                break

        if not stock_col or not date_col:
            result.warnings.append(
                f"Missing required columns. Found: {list(df.columns)}. "
                "Expected at least 'Stock Name' and 'Reference Date'."
            )
            return result

        result.total_rows = len(df)
        today = date.today()
        seen_pairs = set()

        for idx, row in df.iterrows():
            row_num = idx + 2  # 1-indexed plus header
            raw_stock = row[stock_col]
            raw_date = row[date_col]

            # Validate stock string
            stock_str = "" if (raw_stock is None or pd.isna(raw_stock)) else str(raw_stock).strip().upper()
            if not stock_str:
                result.invalid_rows.append(
                    ExcelRowError(row_num, str(raw_stock), str(raw_date), "Stock name is empty.")
                )
                continue

            # Validate reference date
            parsed_dt = self._parse_date(raw_date)
            if not parsed_dt:
                result.invalid_rows.append(
                    ExcelRowError(row_num, stock_str, str(raw_date), "Invalid or unparseable date format.")
                )
                continue

            if parsed_dt > today:
                result.invalid_rows.append(
                    ExcelRowError(row_num, stock_str, str(raw_date), "Reference date cannot be in the future.")
                )
                continue

            # Check duplicate within file
            pair_key = (stock_str, parsed_dt)
            if pair_key in seen_pairs:
                result.warnings.append(f"Row {row_num}: Duplicate stock cycle ({stock_str}, {parsed_dt}) ignored.")
                continue
            seen_pairs.add(pair_key)

            result.valid_rows.append(
                ExcelRowData(
                    row_index=row_num,
                    stock_name=stock_str,
                    reference_date=parsed_dt,
                    raw_date_str=str(raw_date),
                )
            )

        return result

    def import_validated_rows(
        self,
        cycle_service: CycleService,
        valid_rows: List[ExcelRowData],
    ) -> Tuple[int, List[CycleAnalysis]]:
        """Imports valid rows into repository and computes cycle analyses."""
        imported_count = 0
        analyses: List[CycleAnalysis] = []

        for item in valid_rows:
            try:
                _, _, analysis = cycle_service.add_stock_cycle(
                    query=item.stock_name,
                    reference_date=item.reference_date,
                )
                analyses.append(analysis)
                imported_count += 1
            except Exception as e:
                pass

        return imported_count, analyses

    def export_analyses_to_excel(
        self,
        analyses: Sequence[CycleAnalysis],
        output_target: Optional[Union[str, Path, io.BytesIO]] = None,
    ) -> bytes:
        """
        Exports calculated cycle analyses into a complete 18-column Excel spreadsheet.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cycle Analysis"

        # 18 Standard Specification Columns
        headers = [
            "S.No.",
            "Stock Name",
            "Company Name",
            "Cycle",
            "Original Reference Date",
            "Recurring Reference Date",
            "Actual Reference Trading Date",
            "Exchange",
            "Reference High",
            "Current Price",
            "Price Type",
            "Calculation Date",
            "% Change",
            "Bucket",
            "Cycle Start",
            "Cycle End",
            "Data Source",
            "Last Data Refresh",
        ]

        # Header styling
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Segoe UI", size=10)
        border_thin = Side(border_style="thin", color="E2E8F0")
        data_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

        # Write headers
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        ws.row_dimensions[1].height = 28

        # Write data rows
        for idx, item in enumerate(analyses, start=1):
            row_data = [
                idx,
                item.stock_symbol,
                item.company_name,
                f"Cycle {item.cycle_number}",
                item.original_reference_date.strftime("%d-%b-%Y"),
                item.recurring_reference_date.strftime("%d-%b"),
                item.actual_reference_trading_date.strftime("%d-%b-%Y"),
                item.exchange,
                item.reference_high,
                item.current_price,
                item.price_type.value,
                item.calculation_date.strftime("%d-%b-%Y"),
                item.percentage_change / 100.0,
                item.bucket,
                item.cycle_start_date.strftime("%d-%b-%Y"),
                item.cycle_end_date.strftime("%d-%b-%Y"),
                item.data_source,
                item.last_data_refresh.strftime("%Y-%m-%d %H:%M:%S"),
            ]
            ws.append(row_data)
            row_idx = idx + 1
            ws.row_dimensions[row_idx].height = 20

            # Formatting
            for col_num in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.font = data_font
                c.border = data_border
                if col_num in (1, 4, 8, 11, 14, 17):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_num in (9, 10):
                    c.number_format = "₹#,##0.00"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif col_num == 13:
                    c.number_format = "+0.00%;-0.00%;0.00%"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        # Auto column width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Output
        bio = io.BytesIO()
        wb.save(bio)
        content_bytes = bio.getvalue()

        if output_target:
            if isinstance(output_target, (str, Path)):
                Path(output_target).write_bytes(content_bytes)
            elif isinstance(output_target, io.BytesIO):
                output_target.write(content_bytes)

        return content_bytes
